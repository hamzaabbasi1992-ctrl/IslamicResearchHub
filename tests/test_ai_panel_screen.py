"""Tests for the desktop app's collapsible AI-assistant panel."""

import sqlite3
from pathlib import Path

import pytest

pytest.importorskip("PySide6")

from PySide6.QtCore import QSettings  # noqa: E402

from islamic_research_hub.domain.models.book import Book, Page  # noqa: E402
from islamic_research_hub.infrastructure.persistence.master_book_repository import (  # noqa: E402
    MasterBookRepository,
)
from islamic_research_hub.infrastructure.persistence.migration_runner import (  # noqa: E402
    MigrationRunner,
)
from islamic_research_hub.interfaces.desktop_app.ai_panel_screen import (  # noqa: E402
    COLLAPSED_KEY,
    AiAssistantPanel,
)
from islamic_research_hub.interfaces.desktop_app.i18n import Translator  # noqa: E402


def _isolated_settings(tmp_path: Path) -> QSettings:
    return QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat)


def _translator(tmp_path: Path) -> Translator:
    return Translator(_isolated_settings(tmp_path))


def _migrated_database(tmp_path: Path) -> Path:
    """A real, fully-migrated database - content doesn't matter here,
    only that SavedConversations exists so the repository isn't stuck
    in its honest pre-migration no-op mode."""
    database_path = tmp_path / "books.db"
    book = Book(
        information={"Name": "Book of Fiqh"},
        categories=(),
        table_of_contents=(),
        pages=(Page(1, 1, "Content", "Plain"),),
    )
    MasterBookRepository().import_books(
        database_path, (book,), (database_path.parent / "one.mjbz",)
    )
    with sqlite3.connect(database_path) as connection:
        MigrationRunner().migrate(connection)
    return database_path


def test_defaults_to_expanded_when_nothing_is_stored(qtbot, tmp_path: Path) -> None:
    """With no stored preference, the panel starts expanded."""
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path))
    qtbot.addWidget(panel)

    assert panel.is_collapsed is False


def test_toggle_collapsed_flips_state_and_persists(qtbot, tmp_path: Path) -> None:
    """Toggling collapses the panel and remembers it for next time."""
    settings = _isolated_settings(tmp_path)
    panel = AiAssistantPanel(settings, _translator(tmp_path))
    qtbot.addWidget(panel)

    panel.toggle_collapsed()

    assert panel.is_collapsed is True
    assert bool(settings.value(COLLAPSED_KEY, type=bool)) is True


def test_a_second_panel_picks_up_the_persisted_collapsed_state(qtbot, tmp_path: Path) -> None:
    """A fresh panel instance reads back a prior panel's persisted choice."""
    settings = _isolated_settings(tmp_path)
    AiAssistantPanel(settings, _translator(tmp_path)).set_collapsed(True)

    reloaded = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path))
    qtbot.addWidget(reloaded)

    assert reloaded.is_collapsed is True


def test_setting_the_same_state_again_does_not_emit(qtbot, tmp_path: Path) -> None:
    """No redundant signal when the collapsed state doesn't actually change."""
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path))
    qtbot.addWidget(panel)
    received = []
    panel.collapsed_changed.connect(received.append)

    panel.set_collapsed(False)  # already expanded - no-op

    assert received == []


def test_collapsed_changed_signal_carries_the_new_state(qtbot, tmp_path: Path) -> None:
    """The signal argument matches the panel's new collapsed state."""
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path))
    qtbot.addWidget(panel)

    with qtbot.waitSignal(panel.collapsed_changed, timeout=1000) as blocker:
        panel.toggle_collapsed()

    assert blocker.args == [True]


def test_question_input_is_disabled_by_default(qtbot, tmp_path: Path) -> None:
    """Off by default - the AI Agent is a real feature now, but making a
    paid external API call is never silent (same opt-in reasoning as
    TTS/voice search)."""
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path))
    qtbot.addWidget(panel)

    assert not panel._question_edit.isEnabled()
    assert not panel._ask_button.isEnabled()


def test_question_input_is_enabled_when_ai_agent_is_on(qtbot, tmp_path: Path) -> None:
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)

    assert panel._question_edit.isEnabled()
    assert panel._ask_button.isEnabled()


def test_notes_and_references_are_honest_placeholders(qtbot, tmp_path: Path) -> None:
    """Reader Redesign: Notes/References sections are real, visible slots in
    the panel's shape, clearly labeled 'coming soon' - no backend exists for
    either, so nothing here is faked as functional."""
    from PySide6.QtWidgets import QLabel

    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path))
    qtbot.addWidget(panel)

    labels = {label.text() for label in panel.findChildren(QLabel)}
    assert "Notes" in labels
    assert "References" in labels
    assert any("coming soon" in text for text in labels)


class _FakeAiAgentService:
    """A real-shaped, controllable stand-in for AiAgentService."""

    def __init__(self, answer: str = "A grounded answer.", tool_calls_made=("search_books",), fail: bool = False):
        self.answer = answer
        self.tool_calls_made = tool_calls_made
        self.fail = fail
        self.last_question: str | None = None

    def converse(self, question: str):
        from islamic_research_hub.application.ai_agent_service import AgentTurnResult

        self.last_question = question
        if self.fail:
            raise RuntimeError("real failure")
        return AgentTurnResult(answer=self.answer, tool_calls_made=self.tool_calls_made)

    def compare_positions(self, question: str):
        from islamic_research_hub.application.ai_agent_service import AgentTurnResult

        self.last_question = question
        if self.fail:
            raise RuntimeError("real failure")
        return AgentTurnResult(answer=self.answer, tool_calls_made=self.tool_calls_made)


def _install_fake_ai_agent(panel: AiAssistantPanel, **kwargs) -> _FakeAiAgentService:
    """Inject a fake service in place of the real provider-backed one,
    mirroring how test_search_screen.py/test_viewer_screen.py monkeypatch
    their own `_build_real_*_service` seams."""
    fake = _FakeAiAgentService(**kwargs)
    panel._build_real_ai_agent_service = lambda: fake
    return fake


def test_asking_a_question_shows_the_real_answer_and_tool_calls(qtbot, tmp_path: Path) -> None:
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)
    fake = _install_fake_ai_agent(panel, answer="Real grounded answer.", tool_calls_made=("search_books", "get_book_pages"))
    panel._question_edit.setText("What does this library say about patience?")

    panel._on_ask_clicked()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)

    assert fake.last_question == "What does this library say about patience?"
    assert panel._answer_area.toPlainText() == "Real grounded answer."
    assert panel._answer_area.isHidden() is False
    assert "search_books" in panel._tool_calls_label.text()
    assert "get_book_pages" in panel._tool_calls_label.text()
    assert panel._ask_button.isEnabled()


def test_export_answer_button_hidden_until_a_real_answer_arrives(qtbot, tmp_path: Path) -> None:
    """Phase 16 Milestone 1: exporting only makes sense once there's a
    real answer on screen - hidden before that, and hidden again the
    moment a new question is asked."""
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)
    assert panel._export_answer_button.isHidden()
    _install_fake_ai_agent(panel, answer="Real grounded answer.")
    panel._question_edit.setText("A question")

    panel._on_ask_clicked()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)

    assert not panel._export_answer_button.isHidden()

    panel._question_edit.setText("A second question")
    panel._on_ask_clicked()

    assert panel._export_answer_button.isHidden()


def test_export_answer_button_writes_a_real_docx(qtbot, tmp_path: Path, monkeypatch) -> None:
    from docx import Document
    from PySide6.QtWidgets import QFileDialog, QMessageBox

    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)
    _install_fake_ai_agent(panel, answer="Real grounded answer with a citation.")
    panel._question_edit.setText("What does this library say about patience?")
    panel._on_ask_clicked()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)
    output_path = tmp_path / "answer.docx"
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(output_path), ""))
    )
    monkeypatch.setattr(QMessageBox, "information", staticmethod(lambda *a, **k: None))

    panel._on_export_answer_clicked()

    assert output_path.is_file()
    document = Document(output_path)
    all_text = "\n".join(p.text for p in document.paragraphs)
    assert "What does this library say about patience?" in all_text
    assert "Real grounded answer with a citation." in all_text


def test_export_answer_button_does_nothing_when_cancelled(qtbot, tmp_path: Path, monkeypatch) -> None:
    from PySide6.QtWidgets import QFileDialog

    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)
    _install_fake_ai_agent(panel, answer="Real grounded answer.")
    panel._question_edit.setText("A question")
    panel._on_ask_clicked()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: ("", "")))

    panel._on_export_answer_clicked()  # no exception, no file written

    assert not (tmp_path / "answer.docx").exists()


def test_ask_public_method_sets_the_question_and_runs_a_real_turn(qtbot, tmp_path: Path) -> None:
    """`ask()` is the public seam other screens (the Search screen's
    quick-ask box) use to start a real conversation here without
    duplicating this panel's own lazy-build/pre-flight/worker logic -
    same real outcome as typing into the box and clicking Ask."""
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)
    fake = _install_fake_ai_agent(panel, answer="Real grounded answer.")

    panel.ask("What does this library say about patience?")
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)

    assert panel._question_edit.text() == "What does this library say about patience?"
    assert fake.last_question == "What does this library say about patience?"
    assert panel._answer_area.toPlainText() == "Real grounded answer."


def test_compare_mode_checkbox_routes_to_compare_positions(qtbot, tmp_path: Path) -> None:
    """When checked, the worker calls compare_positions(), not converse() -
    verified via a fake that only implements one distinctly."""
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)

    class _CompareOnlyFakeService:
        def __init__(self):
            self.compare_called_with = None

        def compare_positions(self, question: str):
            from islamic_research_hub.application.ai_agent_service import AgentTurnResult

            self.compare_called_with = question
            return AgentTurnResult(answer="Position A vs Position B.", tool_calls_made=("semantic_search_books",))

        def converse(self, question: str):
            raise AssertionError("converse() should not be called in compare mode")

    fake = _CompareOnlyFakeService()
    panel._build_real_ai_agent_service = lambda: fake
    panel._compare_mode_checkbox.setChecked(True)
    panel._question_edit.setText("How did the four madhhabs differ on raising the hands in salah?")

    panel._on_ask_clicked()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)

    assert fake.compare_called_with == "How did the four madhhabs differ on raising the hands in salah?"
    assert panel._answer_area.toPlainText() == "Position A vs Position B."


def test_compare_mode_unchecked_still_routes_to_converse(qtbot, tmp_path: Path) -> None:
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)
    fake = _install_fake_ai_agent(panel, answer="A normal answer.")
    assert panel._compare_mode_checkbox.isChecked() is False
    panel._question_edit.setText("A regular question.")

    panel._on_ask_clicked()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)

    assert fake.last_question == "A regular question."
    assert panel._answer_area.toPlainText() == "A normal answer."


def test_asking_a_question_when_the_service_fails_shows_a_real_message(
    qtbot, tmp_path: Path
) -> None:
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)
    _install_fake_ai_agent(panel, fail=True)
    panel._question_edit.setText("A question")

    panel._on_ask_clicked()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)

    assert panel._answer_area.isHidden() is False
    assert panel._answer_area.toPlainText()  # a real, non-empty failure message
    assert panel._ask_button.isEnabled()


def test_asking_with_no_service_available_shows_a_real_message(qtbot, tmp_path: Path, monkeypatch) -> None:
    """AI Agent enabled but no API key set (or the provider extra isn't
    installed) - _build_real_ai_agent_service() returns None - must
    degrade gracefully, never crash. Also a real misconfiguration, so it
    triggers the shared AI-unavailable popup, not just inline text."""
    popup_calls = []
    monkeypatch.setattr(
        "islamic_research_hub.interfaces.desktop_app.ai_panel_screen.show_ai_unavailable_dialog",
        lambda parent, feature_name, reason: popup_calls.append((feature_name, reason)),
    )
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)
    panel._build_real_ai_agent_service = lambda: None
    panel._question_edit.setText("A question")

    panel._on_ask_clicked()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)

    assert panel._answer_area.isHidden() is False
    assert "unavailable" in panel._answer_area.toPlainText().lower()
    assert len(popup_calls) == 1
    assert popup_calls[0][0] == "AI Agent"


def test_a_genuine_runtime_failure_does_not_trigger_the_unavailable_popup(
    qtbot, tmp_path: Path, monkeypatch
) -> None:
    """A mid-request exception (network/API error) is a different, often
    transient case from "not configured" - inline text only, no popup."""
    popup_calls = []
    monkeypatch.setattr(
        "islamic_research_hub.interfaces.desktop_app.ai_panel_screen.show_ai_unavailable_dialog",
        lambda parent, feature_name, reason: popup_calls.append((feature_name, reason)),
    )
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)
    _install_fake_ai_agent(panel, fail=True)
    panel._question_edit.setText("A question")

    panel._on_ask_clicked()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)

    assert panel._answer_area.isHidden() is False
    assert popup_calls == []


def test_ask_button_disabled_while_a_request_is_in_flight(qtbot, tmp_path: Path) -> None:
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)
    _install_fake_ai_agent(panel)
    panel._question_edit.setText("A question")

    panel._on_ask_clicked()

    assert not panel._ask_button.isEnabled()
    assert not panel._question_edit.isEnabled()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass


def test_asking_a_blank_question_does_nothing(qtbot, tmp_path: Path) -> None:
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path), enable_lazy_ai_agent=True)
    qtbot.addWidget(panel)
    _install_fake_ai_agent(panel)

    panel._on_ask_clicked()

    assert panel._ai_agent_worker is None


def test_build_llm_provider_returns_none_for_an_unknown_provider(qtbot, tmp_path: Path) -> None:
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path))
    qtbot.addWidget(panel)

    assert panel._build_llm_provider("not-a-real-provider", "some-key") is None


def test_build_llm_provider_for_ollama_needs_no_real_api_key(qtbot, tmp_path: Path) -> None:
    """Ollama is a real local model - it never uses the `api_key` arg at
    all, unlike the 3 cloud providers."""
    from islamic_research_hub.infrastructure.ai.ollama_llm_provider import OllamaLlmProvider

    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path))
    qtbot.addWidget(panel)

    provider = panel._build_llm_provider("ollama", "ollama-local")

    assert isinstance(provider, OllamaLlmProvider)


def test_build_llm_provider_for_ollama_reads_the_real_model_and_url_from_settings(
    qtbot, tmp_path: Path
) -> None:
    from islamic_research_hub.interfaces.desktop_app.settings_screen import (
        OLLAMA_BASE_URL_KEY,
        OLLAMA_MODEL_KEY,
    )

    settings = _isolated_settings(tmp_path)
    settings.setValue(OLLAMA_MODEL_KEY, "llama3.1")
    settings.setValue(OLLAMA_BASE_URL_KEY, "http://192.168.1.50:11434/v1")
    panel = AiAssistantPanel(settings, _translator(tmp_path))
    qtbot.addWidget(panel)

    provider = panel._build_llm_provider("ollama", "ollama-local")

    assert provider._model == "llama3.1"


def test_lazy_ai_agent_is_not_attempted_by_default(qtbot, tmp_path: Path) -> None:
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path))
    qtbot.addWidget(panel)
    build_calls = []
    panel._build_real_ai_agent_service = lambda: (build_calls.append(1), None)[1]

    result = panel.get_or_build_ai_agent_service()

    assert result is None
    assert build_calls == []


def test_switching_language_retranslates_the_panel(qtbot, tmp_path: Path) -> None:
    settings = _isolated_settings(tmp_path)
    translator = Translator(settings)
    panel = AiAssistantPanel(settings, translator)
    qtbot.addWidget(panel)
    assert panel._title_label.text() == "Assistant"
    assert panel._ask_button.text() == "Ask"

    translator.set_language("ur")

    assert panel._title_label.text() == "اسسٹنٹ"
    assert panel._ask_button.text() == "پوچھیں"
    assert panel._export_answer_button.text() == "جواب برآمد کریں"
    assert panel._notes_heading.text() == "نوٹس"
    assert panel._compare_mode_checkbox.text() == "علمی آراء کا موازنہ کریں"


def test_saved_conversations_button_disabled_without_a_database(qtbot, tmp_path: Path) -> None:
    """No database means no `SavedConversationRepository` to talk to -
    the button is disabled rather than crashing on click."""
    panel = AiAssistantPanel(_isolated_settings(tmp_path), _translator(tmp_path))
    qtbot.addWidget(panel)

    assert panel._saved_conversations is None
    assert not panel._saved_conversations_button.isEnabled()


def test_save_conversation_button_hidden_until_a_real_answer_arrives(qtbot, tmp_path: Path) -> None:
    database_path = _migrated_database(tmp_path)
    panel = AiAssistantPanel(
        _isolated_settings(tmp_path),
        _translator(tmp_path),
        database_path=database_path,
        enable_lazy_ai_agent=True,
    )
    qtbot.addWidget(panel)
    assert panel._save_conversation_button.isHidden()
    _install_fake_ai_agent(panel, answer="Real grounded answer.")
    panel._question_edit.setText("A question")

    panel._on_ask_clicked()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)

    assert not panel._save_conversation_button.isHidden()

    panel._question_edit.setText("A second question")
    panel._on_ask_clicked()

    assert panel._save_conversation_button.isHidden()


def test_saving_a_conversation_persists_the_real_question_and_answer(
    qtbot, tmp_path: Path, monkeypatch
) -> None:
    from islamic_research_hub.infrastructure.persistence.saved_conversation_repository import (
        SavedConversationRepository,
    )

    database_path = _migrated_database(tmp_path)
    panel = AiAssistantPanel(
        _isolated_settings(tmp_path),
        _translator(tmp_path),
        database_path=database_path,
        enable_lazy_ai_agent=True,
    )
    qtbot.addWidget(panel)
    _install_fake_ai_agent(panel, answer="Real grounded answer.")
    panel._question_edit.setText("What does this library say about patience?")
    panel._on_ask_clicked()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)
    monkeypatch.setattr(
        "islamic_research_hub.interfaces.desktop_app.ai_panel_screen.QInputDialog.getText",
        staticmethod(lambda *a, **k: ("Patience answer", True)),
    )

    panel._on_save_conversation_clicked()

    saved = SavedConversationRepository(database_path).list_conversations()
    assert len(saved) == 1
    assert saved[0].name == "Patience answer"
    assert saved[0].question == "What does this library say about patience?"
    assert saved[0].answer == "Real grounded answer."


def test_saving_a_conversation_with_a_taken_name_shows_a_warning(
    qtbot, tmp_path: Path, monkeypatch
) -> None:
    database_path = _migrated_database(tmp_path)
    panel = AiAssistantPanel(
        _isolated_settings(tmp_path),
        _translator(tmp_path),
        database_path=database_path,
        enable_lazy_ai_agent=True,
    )
    qtbot.addWidget(panel)
    panel._saved_conversations.save_conversation("Taken name", "Question one", "Answer one")
    _install_fake_ai_agent(panel, answer="Real grounded answer.")
    panel._question_edit.setText("A question")
    panel._on_ask_clicked()
    with qtbot.waitSignal(panel._ai_agent_worker.finished, timeout=5000):
        pass
    qtbot.wait(50)
    monkeypatch.setattr(
        "islamic_research_hub.interfaces.desktop_app.ai_panel_screen.QInputDialog.getText",
        staticmethod(lambda *a, **k: ("Taken name", True)),
    )
    warning_calls = []
    monkeypatch.setattr(
        "islamic_research_hub.interfaces.desktop_app.ai_panel_screen.QMessageBox.warning",
        staticmethod(lambda *a, **k: warning_calls.append(a)),
    )

    panel._on_save_conversation_clicked()

    assert len(warning_calls) == 1


def test_saved_conversations_dialog_shows_the_real_empty_state(
    qtbot, tmp_path: Path, monkeypatch
) -> None:
    from PySide6.QtWidgets import QLabel

    database_path = _migrated_database(tmp_path)
    panel = AiAssistantPanel(
        _isolated_settings(tmp_path),
        _translator(tmp_path),
        database_path=database_path,
        enable_lazy_ai_agent=True,
    )
    qtbot.addWidget(panel)
    shown_dialogs = []
    monkeypatch.setattr(
        "islamic_research_hub.interfaces.desktop_app.ai_panel_screen.QDialog.exec",
        lambda self: shown_dialogs.append(self) or None,
    )

    panel._on_view_saved_conversations_clicked()

    assert len(shown_dialogs) == 1
    labels = [label.text() for label in shown_dialogs[0].findChildren(QLabel)]
    assert any("No saved conversations yet" in text for text in labels)


def test_saved_conversations_dialog_opens_a_real_saved_conversation(
    qtbot, tmp_path: Path, monkeypatch
) -> None:
    from PySide6.QtWidgets import QPushButton

    database_path = _migrated_database(tmp_path)
    panel = AiAssistantPanel(
        _isolated_settings(tmp_path),
        _translator(tmp_path),
        database_path=database_path,
        enable_lazy_ai_agent=True,
    )
    qtbot.addWidget(panel)
    panel._saved_conversations.save_conversation(
        "Patience answer", "What does this library say about patience?", "A saved answer."
    )
    shown_dialogs = []
    monkeypatch.setattr(
        "islamic_research_hub.interfaces.desktop_app.ai_panel_screen.QDialog.exec",
        lambda self: shown_dialogs.append(self) or None,
    )
    panel._on_view_saved_conversations_clicked()
    open_button = next(
        b for b in shown_dialogs[0].findChildren(QPushButton) if b.text() == "Run"
    )

    open_button.click()

    assert panel._question_edit.text() == "What does this library say about patience?"
    assert panel._answer_area.toPlainText() == "A saved answer."
    assert not panel._answer_area.isHidden()


def test_deleting_a_saved_conversation_from_the_dialog_removes_it(
    qtbot, tmp_path: Path, monkeypatch
) -> None:
    from PySide6.QtWidgets import QPushButton

    database_path = _migrated_database(tmp_path)
    panel = AiAssistantPanel(
        _isolated_settings(tmp_path),
        _translator(tmp_path),
        database_path=database_path,
        enable_lazy_ai_agent=True,
    )
    qtbot.addWidget(panel)
    panel._saved_conversations.save_conversation("Patience answer", "Question", "Answer")
    shown_dialogs = []
    monkeypatch.setattr(
        "islamic_research_hub.interfaces.desktop_app.ai_panel_screen.QDialog.exec",
        lambda self: shown_dialogs.append(self) or None,
    )
    panel._on_view_saved_conversations_clicked()
    delete_button = next(
        b for b in shown_dialogs[0].findChildren(QPushButton) if b.text() == "Delete"
    )

    delete_button.click()

    assert panel._saved_conversations.list_conversations() == ()
